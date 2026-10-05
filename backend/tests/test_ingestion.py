import tempfile
import uuid
from pathlib import Path

import docx
import pymupdf
import pytest

from app.rag.ingestion.base import (
    DocumentParsingError,
    EmptyDocumentError,
    UnsupportedFileTypeError,
)
from app.rag.ingestion.docx import DocxParser
from app.rag.ingestion.markdown import MarkdownParser
from app.rag.ingestion.pdf import PDFParser
from app.rag.ingestion.service import IngestionService, ingestion_service
from app.rag.ingestion.txt import TextParser


# ─── PDF Parser Tests ──────────────────────────────────────────────
def test_pdf_parser_success(tmp_path: Path):
    pdf_path = tmp_path / "sample.pdf"
    doc = pymupdf.open()
    
    # Page 1
    page1 = doc.new_page()
    page1.insert_text((50, 50), "Hello PDF Page 1", fontsize=12)
    
    # Page 2
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Second Page Content", fontsize=12)
    
    doc.save(str(pdf_path))
    doc.close()

    parser = PDFParser()
    doc_id = uuid.uuid4()
    result = parser.parse(pdf_path, doc_id)

    assert result["document_id"] == str(doc_id)
    assert result["source"] == "sample.pdf"
    assert result["file_type"] == "pdf"
    assert result["metadata"]["total_pages"] == 2
    assert len(result["pages"]) == 2
    assert "Hello PDF Page 1" in result["full_text"]
    assert "Second Page Content" in result["full_text"]


def test_pdf_parser_nonexistent_file():
    parser = PDFParser()
    with pytest.raises(DocumentParsingError):
        parser.parse(Path("/nonexistent/file.pdf"), uuid.uuid4())


# ─── Markdown Parser Tests ──────────────────────────────────────────
def test_markdown_parser_success(tmp_path: Path):
    md_path = tmp_path / "sample.md"
    md_content = """# Title Section

This is paragraph 1.

- Item 1
- Item 2

## Subsection

```python
def hello():
    return "world"
```
"""
    md_path.write_text(md_content, encoding="utf-8")

    parser = MarkdownParser()
    doc_id = uuid.uuid4()
    result = parser.parse(md_path, doc_id)

    assert result["document_id"] == str(doc_id)
    assert result["file_type"] == "markdown"
    assert len(result["sections"]) >= 2
    assert result["sections"][0]["heading"] == "Title Section"
    assert "hello():" in result["full_text"]


def test_markdown_parser_empty_file(tmp_path: Path):
    empty_md = tmp_path / "empty.md"
    empty_md.write_text("", encoding="utf-8")

    parser = MarkdownParser()
    with pytest.raises(EmptyDocumentError):
        parser.parse(empty_md, uuid.uuid4())


# ─── DOCX Parser Tests ─────────────────────────────────────────────
def test_docx_parser_success(tmp_path: Path):
    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Document Title", level=1)
    doc.add_paragraph("Introductory paragraph text.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(1, 0).text = "Val 1"
    table.cell(1, 1).text = "Val 2"
    
    doc.save(str(docx_path))

    parser = DocxParser()
    doc_id = uuid.uuid4()
    result = parser.parse(docx_path, doc_id)

    assert result["document_id"] == str(doc_id)
    assert result["file_type"] == "docx"
    assert result["metadata"]["table_count"] == 1
    assert "Introductory paragraph text." in result["full_text"]
    assert "Val 1" in result["full_text"]


# ─── Text Parser Tests ─────────────────────────────────────────────
def test_text_parser_success(tmp_path: Path):
    txt_path = tmp_path / "notes.txt"
    txt_path.write_text("Line 1 text.\n\nLine 2 text.", encoding="utf-8")

    parser = TextParser()
    doc_id = uuid.uuid4()
    result = parser.parse(txt_path, doc_id)

    assert result["file_type"] == "txt"
    assert result["metadata"]["paragraph_count"] == 2
    assert "Line 1 text." in result["full_text"]


# ─── Ingestion Service Tests ───────────────────────────────────────
def test_ingestion_service_routing(tmp_path: Path):
    service = IngestionService()

    # Create dummy pdf
    pdf_file = tmp_path / "doc.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "PDF content", fontsize=12)
    doc.save(str(pdf_file))
    doc.close()

    result = service.parse(pdf_file, uuid.uuid4())
    assert result["file_type"] == "pdf"


def test_ingestion_service_unsupported_file_type(tmp_path: Path):
    unsupported_file = tmp_path / "file.unknown"
    unsupported_file.write_text("data", encoding="utf-8")

    service = IngestionService()
    with pytest.raises(UnsupportedFileTypeError):
        service.parse(unsupported_file, uuid.uuid4())
